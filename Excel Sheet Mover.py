# This program copies a sheet from one workbook to another, preserving formatting and values, but not formulas.
# It works with Excel 2010 and later. It uses openpyxl library to handle Excel files.

import openpyxl
from openpyxl import load_workbook

import tkinter as tk
from tkinter import filedialog, messagebox
import copy

from pathlib import Path

# Function to copy a sheet from source workbook to destination workbook
def copy_sheet(source_file, destination_file, sheet_name):
    destination_file = Path(destination_file)

    # Load the source workbook once for formatting and once for cached values.
    source_wb = load_workbook(source_file, data_only=False)
    source_values_wb = load_workbook(source_file, data_only=True)
    if sheet_name not in source_wb.sheetnames:
        messagebox.showerror("Error", f"Sheet '{sheet_name}' not found in source workbook.")
        return
    
    source_sheet = source_wb[sheet_name]
    source_values_sheet = source_values_wb[sheet_name]

    # Load or create the destination workbook
    if destination_file.exists():
        destination_wb = load_workbook(destination_file)
    else:
        destination_wb = openpyxl.Workbook()
        destination_wb.remove(destination_wb.active)  # Remove the default sheet

    # Create a new sheet in the destination workbook with the same name
    if sheet_name in destination_wb.sheetnames:
        messagebox.showerror("Error", f"Sheet '{sheet_name}' already exists in destination workbook.")
        return
    
    destination_sheet = destination_wb.create_sheet(title=sheet_name)

    # Copy cell values and formatting from source to destination
    for row in source_sheet.iter_rows():
        for cell in row:
            value = source_values_sheet.cell(row=cell.row, column=cell.column).value
            new_cell = destination_sheet.cell(row=cell.row, column=cell.column, value=value)
            if cell.has_style:
                new_cell.font = copy.copy(cell.font)
                new_cell.border = copy.copy(cell.border)
                new_cell.fill = copy.copy(cell.fill)
                new_cell.number_format = copy.copy(cell.number_format)
                new_cell.protection = copy.copy(cell.protection)
                new_cell.alignment = copy.copy(cell.alignment)

    # Copy merged cell ranges, such as A1:C1.
    for merged_range in source_sheet.merged_cells.ranges:
        destination_sheet.merge_cells(str(merged_range))

    # copy column widths
    try:
        for col, dim in source_sheet.column_dimensions.items():
            if dim and dim.width:
                destination_sheet.column_dimensions[col].width = dim.width
    except Exception:
        pass

    # copy row heights
    try:
        for idx, dim in source_sheet.row_dimensions.items():
            if dim and dim.height:
                destination_sheet.row_dimensions[idx].height = dim.height
    except Exception:
        pass
    
    # Save the destination workbook
    destination_wb.save(destination_file)
    messagebox.showinfo("Success", f"Sheet '{sheet_name}' copied successfully to '{destination_file}'.")

# GUI setup to allow user to select source and destination files and specify the sheet name
def main():
    root = tk.Tk()
    root.withdraw()  # Hide the main window

    # Ask user to select source workbook
    source_file = filedialog.askopenfilename(title="Select Source Workbook", filetypes=[("Excel files", "*.xlsx")])
    if not source_file:
        messagebox.showerror("Error", "No source workbook selected.")
        return

    # Ask user to select destination workbook
    destination_file = filedialog.askopenfilename(title="Select Destination Workbook", filetypes=[("Excel files", "*.xlsx")])
    if not destination_file:
        messagebox.showerror("Error", "No destination workbook selected.")
        return

    # Ask user for the sheet name to copy
    sheet_name = tk.simpledialog.askstring("Input", "Enter the name of the sheet to copy:")
    if not sheet_name:
        messagebox.showerror("Error", "No sheet name provided.")
        return

    # Call the function to copy the sheet
    copy_sheet(source_file, destination_file, sheet_name)

#Run program
main()